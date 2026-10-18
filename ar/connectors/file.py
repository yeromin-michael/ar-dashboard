"""CSV/XLSX connector driven by a column-mapping JSON file.

The mapping says, per record type, which file to read and which source column
feeds each canonical field. Only `invoices` is required; customers are derived
from invoice rows when no customers file is given.

Mapping shape (see mappings/example.json):
{
  "date_formats": ["%d/%m/%Y", "%Y-%m-%d"],
  "files": {
    "invoices":  {"path": "invoices.csv", "sheet": null,
                  "columns": {"id": "Invoice No", "customer_id": "Customer", ...},
                  "defaults": {"currency": "USD", "entity": "Inc"}},
    "customers": {...}, "payments": {...}, "revenue": {...},
    "contracts": {...}, "usage": {...}, "actions": {...}
  }
}
"""
from __future__ import annotations

import csv
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from ..model import (ActionItem, ContractTerm, Customer, Dataset, Invoice,
                     Payment, RevenueLine, UsageBatch)
from ..profile import due_date

DEFAULT_DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%m/%d/%Y", "%d.%m.%Y", "%d %b %y", "%d %b %Y"]
CURRENCY_SIGNS = {"$": "USD", "€": "EUR", "₪": "ILS"}


class MappingError(ValueError):
    pass


# ---------- reading ----------

def read_rows(path: Path, sheet: str | None = None) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise MappingError("reading .xlsx needs openpyxl (pip install openpyxl)") from e
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        header_idx = next((i for i, r in enumerate(rows) if sum(1 for v in r if v not in (None, "")) >= 2), None)
        if header_idx is None:
            return []
        header = [str(h).strip() if h is not None else "" for h in rows[header_idx]]
        return [dict(zip(header, r)) for r in rows[header_idx + 1:] if any(v not in (None, "") for v in r)]
    with path.open(newline="", encoding="utf-8-sig") as fh:
        return [{(k or "").strip(): v for k, v in row.items()} for row in csv.DictReader(fh)
                if any((v or "").strip() for v in row.values() if isinstance(v, str))]


# ---------- parsing ----------

def parse_amount(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip()
    neg = t.startswith("(") and t.endswith(")")
    t = re.sub(r"[^\d.,\-]", "", t)
    if t.count(",") and t.count("."):
        t = t.replace(",", "") if t.rfind(".") > t.rfind(",") else t.replace(".", "").replace(",", ".")
    elif t.count(",") == 1 and len(t.split(",")[1]) in (1, 2):
        t = t.replace(",", ".")               # decimal comma
    else:
        t = t.replace(",", "")
    if t in ("", "-", "."):
        raise ValueError(f"not an amount: {v!r}")
    x = float(t)
    return -abs(x) if neg else x


def parse_date(v, formats: list[str]) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):            # Excel serial
        return date(1899, 12, 30) + timedelta(days=int(v))
    t = str(v).strip()
    for f in formats:
        try:
            return datetime.strptime(t, f).date()
        except ValueError:
            continue
    raise ValueError(f"unrecognized date {v!r}; add its format to date_formats")


def parse_currency(v, default: str) -> str:
    t = str(v or "").strip().upper()
    if not t:
        return default
    return CURRENCY_SIGNS.get(t, t)


def slug(t: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", str(t).lower()).strip("-")


def parse_timing(v) -> str:
    """Map export wording to a revenue type; empty = derive from category."""
    t = re.sub(r"[^a-z]", "", str(v or "").lower())
    if not t:
        return ""
    if t in ("overtime", "ratably", "ratable", "ot"):
        return "over_time"
    if t in ("pointintime", "atapointintime", "pit", "pointin", "upondelivery"):
        return "point_in_time"
    raise ValueError(f"unknown revenue type {v!r} (use 'over time' or 'point in time')")


def _list(v) -> list[str]:
    return [x.strip() for x in re.split(r"[;,]", str(v or "")) if x.strip()]


# ---------- mapping ----------

class Mapper:
    def __init__(self, spec: dict, kind: str, formats: list[str]):
        self.kind, self.formats = kind, formats
        self.columns = spec.get("columns", {})
        self.defaults = spec.get("defaults", {})

    def raw(self, row: dict, field: str, required: bool = False):
        col = self.columns.get(field)
        val = row.get(col) if col else None
        if (val is None or val == "") and field in self.defaults:
            val = self.defaults[field]
        if required and (val is None or val == ""):
            src = f"column '{col}'" if col else "no column mapped"
            raise MappingError(f"{self.kind}: missing required field '{field}' ({src})")
        return val

    def text(self, row, field, required=False, default=""):
        v = self.raw(row, field, required)
        return default if v is None or v == "" else str(v).strip()

    def amount(self, row, field, required=False):
        return parse_amount(self.raw(row, field, required))

    def date(self, row, field, required=False):
        v = self.raw(row, field, required)
        return None if v is None or v == "" else parse_date(v, self.formats)


def load(options: dict, as_of: date) -> Dataset:
    base = Path(options.get("dir", "data/private"))
    mapping_path = Path(options.get("mapping", "mappings/example.json"))
    spec = json.loads(mapping_path.read_text(encoding="utf-8"))
    return load_mapping(spec, base, source=f"file ({base.name})", terms=options.get("terms"))


def load_mapping(spec: dict, base: Path, source: str = "file", terms: dict | None = None) -> Dataset:
    """terms = the profile's payment_terms; used when a row has no due date or terms."""
    default_days = int((terms or {}).get("default_days", 30))
    terms_basis = (terms or {}).get("basis", "net")
    formats = spec.get("date_formats") or DEFAULT_DATE_FORMATS
    files = spec.get("files", {})
    if "invoices" not in files:
        raise MappingError("mapping must define files.invoices")
    ds = Dataset(source=source)

    def rows_of(kind):
        f = files.get(kind)
        if not f:
            return None, []
        path = base / f["path"]
        if not path.exists():
            raise MappingError(f"{kind}: file not found: {path}")
        return Mapper(f, kind, formats), read_rows(path, f.get("sheet"))

    def err(kind, i, e):
        return MappingError(f"{kind} row {i + 2}: {e}")   # +2: header row, 1-based

    # customers (optional)
    m, rows = rows_of("customers")
    for i, r in enumerate(rows):
        try:
            name = m.text(r, "name", required=True)
            ds.customers.append(Customer(
                id=m.text(r, "id") or slug(name), name=name,
                entity=m.text(r, "entity", default="Inc"),
                currency=parse_currency(m.raw(r, "currency"), "USD"),
                product=m.text(r, "product"), contact_name=m.text(r, "contact_name"),
                contact_emails=_list(m.raw(r, "contact_emails")),
                budget_code=m.text(r, "budget_code"),
                terms_days=int(m.amount(r, "terms_days") or default_days),
                modules=_list(m.raw(r, "modules")),
            ))
        except (ValueError, MappingError) as e:
            raise err("customers", i, e)
    known = {c.id for c in ds.customers}
    by_name = {c.name.lower(): c.id for c in ds.customers}

    def customer_ref(m: Mapper, r: dict, currency: str, entity: str) -> str:
        cid = m.text(r, "customer_id")
        name = m.text(r, "customer_name")
        if not cid and name:
            cid = by_name.get(name.lower()) or slug(name)
        if not cid:
            raise MappingError("needs customer_id or customer_name")
        if cid not in known:
            if files.get("customers"):
                raise MappingError(f"unknown customer '{cid}'")
            ds.customers.append(Customer(id=cid, name=name or cid, entity=entity, currency=currency,
                                         terms_days=default_days))
            known.add(cid)
            by_name[(name or cid).lower()] = cid
        return cid

    # invoices
    m, rows = rows_of("invoices")
    for i, r in enumerate(rows):
        try:
            cur = parse_currency(m.raw(r, "currency"), "USD")
            entity = m.text(r, "entity", default="Inc")
            issue = m.date(r, "issue_date", required=True)
            due = m.date(r, "due_date") or due_date(issue, int(m.amount(r, "terms_days") or default_days), terms_basis)
            if m.columns.get("net") or "net" in m.defaults:
                net, tax = m.amount(r, "net", True), m.amount(r, "tax")
            else:   # only a gross total available
                total, tax = m.amount(r, "total", True), m.amount(r, "tax")
                net = round(total - tax, 2)
            ds.invoices.append(Invoice(
                id=m.text(r, "id", required=True), customer_id=customer_ref(m, r, cur, entity), entity=entity,
                issue_date=issue, due_date=due, currency=cur, net=net, tax=tax,
                category=m.text(r, "category", default="subscription").lower(),
                kind=m.text(r, "kind", default="new").lower(),
                description=m.text(r, "description"), pdf_url=m.text(r, "pdf_url"),
            ))
        except (ValueError, MappingError) as e:
            raise err("invoices", i, e)
    inv_by_id = {i.id: i for i in ds.invoices}

    # payments (optional). Amounts are stored positive = reduces balance;
    # set "negate": true in the mapping when the export shows receipts as negatives.
    m, rows = rows_of("payments")
    negate = bool(files.get("payments", {}).get("negate"))
    for i, r in enumerate(rows):
        try:
            inv_id = m.text(r, "invoice_id") or None
            inv = inv_by_id.get(inv_id) if inv_id else None
            cur = parse_currency(m.raw(r, "currency"), inv.currency if inv else "USD")
            amount = m.amount(r, "amount", True)
            ds.payments.append(Payment(
                id=m.text(r, "id") or f"pay-{i + 1}",
                customer_id=inv.customer_id if inv else customer_ref(m, r, cur, "Inc"),
                date=m.date(r, "date", required=True), currency=cur,
                amount=-amount if negate else amount,
                kind=m.text(r, "kind", default="payment").lower(),
                invoice_id=inv_id, reference=m.text(r, "reference"),
            ))
        except (ValueError, MappingError) as e:
            raise err("payments", i, e)

    # revenue (optional)
    m, rows = rows_of("revenue")
    for i, r in enumerate(rows):
        try:
            month = m.text(r, "month", required=True)
            if not re.fullmatch(r"\d{4}-\d{2}", month):
                month = parse_date(month, formats).strftime("%Y-%m")
            cur = parse_currency(m.raw(r, "currency"), "USD")
            ds.revenue.append(RevenueLine(customer_ref(m, r, cur, "Inc"), month,
                                          m.text(r, "category", default="subscription").lower(),
                                          m.amount(r, "amount", True), cur, m.text(r, "invoice_id") or None,
                                          parse_timing(m.text(r, "timing"))))
        except (ValueError, MappingError) as e:
            raise err("revenue", i, e)

    # contracts / usage / actions (optional)
    m, rows = rows_of("contracts")
    for i, r in enumerate(rows):
        try:
            ds.contracts.append(ContractTerm(customer_ref(m, r, "USD", "Inc"), m.text(r, "record_type", True),
                                             m.text(r, "description"), m.text(r, "key_fields")))
        except (ValueError, MappingError) as e:
            raise err("contracts", i, e)
    m, rows = rows_of("usage")
    for i, r in enumerate(rows):
        try:
            ds.usage.append(UsageBatch(customer_ref(m, r, "USD", "Inc"), m.text(r, "month", True),
                                       int(m.amount(r, "units", True)), m.amount(r, "rate_per_unit", True),
                                       m.amount(r, "base_fee"), m.amount(r, "tax_rate"), m.amount(r, "hw_share")))
        except (ValueError, MappingError) as e:
            raise err("usage", i, e)
    m, rows = rows_of("actions")
    for i, r in enumerate(rows):
        try:
            ds.actions.append(ActionItem(m.text(r, "id") or f"act-{i + 1}", customer_ref(m, r, "USD", "Inc"),
                                         m.text(r, "kind", default="action").lower(), m.text(r, "title", True),
                                         m.text(r, "body"), m.text(r, "urgency", default="medium").lower(),
                                         m.text(r, "tag"), m.date(r, "date")))
        except (ValueError, MappingError) as e:
            raise err("actions", i, e)

    # derive revenue from invoices when no revenue file is given
    if not files.get("revenue"):
        for inv in ds.invoices:
            ds.revenue.append(RevenueLine(inv.customer_id, inv.issue_date.strftime("%Y-%m"),
                                          inv.category, inv.net, inv.currency, inv.id))
    return ds
